import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from pathlib import Path

def convert_md_to_docx(md_path: Path, docx_path: Path):
    content = md_path.read_text(encoding="utf-8")
    doc = Document()

    # Set 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(20)
        run.font.bold = True
        run.font.color.rgb = RGBColor(16, 52, 96)

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 76, 129)

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(40, 40, 40)

    def add_heading_3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(70, 70, 70)

    def add_body(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = "Arial"
            r_bold.font.size = Pt(10)
            r_bold.font.bold = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10)

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = "Arial"
            r_bold.font.size = Pt(10)
            r_bold.font.bold = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10)

    def add_code_block(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.font.name = "Courier New"
        run.font.size = Pt(9)
        shd_xml = f'<w:shd {nsdecls("w")} w:fill="F4F5F7"/>'
        p._p.get_or_add_pPr().append(parse_xml(shd_xml))

    lines = content.split("\n")
    in_code = False
    code_lines = []

    for line in lines:
        raw = line.rstrip()
        if raw.startswith("```"):
            if in_code:
                add_code_block("\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
                code_lines = []
            continue

        if in_code:
            code_lines.append(raw)
            continue

        if raw.startswith("# "):
            add_title(raw[2:].strip())
        elif raw.startswith("## "):
            add_heading_1(raw[3:].strip())
        elif raw.startswith("### "):
            add_heading_2(raw[4:].strip())
        elif raw.startswith("#### "):
            add_heading_3(raw[5:].strip())
        elif raw.startswith("- ") or raw.startswith("* "):
            item_text = raw[2:].strip()
            if "**" in item_text:
                parts = item_text.split("**")
                if len(parts) >= 3 and item_text.startswith("**"):
                    add_bullet(item_text[len(parts[1])+4:], bold_prefix=parts[1] + ": ")
                else:
                    add_bullet(item_text.replace("**", ""))
            else:
                add_bullet(item_text)
        elif raw.startswith("> "):
            add_body(raw[2:].strip())
        elif raw.strip() == "---":
            pass
        elif raw.strip():
            clean_line = raw.strip()
            if clean_line.startswith("|"):
                add_body(clean_line)
            else:
                add_body(clean_line.replace("**", ""))

    doc.save(str(docx_path))
    print(f"Successfully generated {docx_path} ({docx_path.stat().st_size} bytes)")

if __name__ == "__main__":
    md_file = Path("docs/HERMES_PROJECT_EXPLANATION.md")
    docx_file = Path("docs/HERMES_PROJECT_EXPLANATION.docx")
    convert_md_to_docx(md_file, docx_file)
